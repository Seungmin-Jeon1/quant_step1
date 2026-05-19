from __future__ import annotations

import subprocess
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd


PROJECT_ROOT = Path(__file__).resolve().parent
OUTPUT_DIR = PROJECT_ROOT / "output"
FINAL_REPORT_XLSX = OUTPUT_DIR / "final_daily_research_report.xlsx"

PIPELINE_STEPS = [
    ("Current screener", PROJECT_ROOT / "screening" / "screen_stocks.py", True),
    ("Signal backtest", PROJECT_ROOT / "backtesting" / "backtest_growth.py", True),
    ("Backtest interpretation", PROJECT_ROOT / "backtesting" / "analyze_backtest.py", True),
    ("Enhanced backtest", PROJECT_ROOT / "backtesting" / "enhance_backtest_growth.py", True),
    (
        "First-hit backtest",
        PROJECT_ROOT / "backtesting" / "first_hit_backtest_growth.py",
        False,
    ),
    ("Daily priority ranking", PROJECT_ROOT / "rank_screening_results.py", True),
    ("Growth risk score", PROJECT_ROOT / "risk" / "risk_score_growth.py", True),
    ("Long-term quality score", PROJECT_ROOT / "long_term_quality_score.py", True),
    ("Event risk check", PROJECT_ROOT / "event_risk_checker.py", True),
    ("Position alerts", PROJECT_ROOT / "position_alert_checker.py", False),
]

PRIORITY_CSV = OUTPUT_DIR / "daily_growth_priority.csv"
RISK_CSV = OUTPUT_DIR / "growth_risk_score.csv"
QUALITY_CSV = OUTPUT_DIR / "long_term_growth_quality.csv"
EVENT_CSV = OUTPUT_DIR / "growth_event_risk.csv"
POSITION_ALERTS_CSV = OUTPUT_DIR / "position_alerts.csv"
MEGA_CAP_VALUATION_CSV = (
    PROJECT_ROOT / "screening" / "output" / "mega_cap_valuation_screen.csv"
)
ENHANCED_SUMMARY_CSV = (
    PROJECT_ROOT / "backtesting" / "output" / "growth_signal_enhanced_backtest_summary.csv"
)
FIRST_HIT_SUMMARY_CSV = (
    PROJECT_ROOT / "backtesting" / "output" / "growth_signal_first_hit_summary.csv"
)
FIRST_HIT_BY_TICKER_CSV = (
    PROJECT_ROOT / "backtesting" / "output" / "growth_signal_first_hit_by_ticker.csv"
)
INTERPRETATION_CSV = (
    PROJECT_ROOT / "backtesting" / "output" / "growth_signal_backtest_interpretation.csv"
)


def text_value(value: Any) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except TypeError:
        pass
    return str(value)


def number_value(value: Any) -> float | None:
    try:
        if pd.isna(value):
            return None
    except TypeError:
        pass
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def parse_note_fields(note: Any) -> dict[str, str]:
    fields: dict[str, str] = {}
    for part in text_value(note).split(";"):
        if "=" not in part:
            continue
        key, value = part.split("=", 1)
        fields[key.strip()] = value.strip() or "확인 필요"
    return fields


def korean_interpretation_for_row(row: dict[str, Any]) -> str:
    metric = text_value(row.get("metric"))
    value = text_value(row.get("value")) or "확인 필요"
    note = text_value(row.get("note"))
    note_fields = parse_note_fields(note)

    if metric == "report_generated_at":
        return (
            "이 리포트가 생성된 시각입니다. 매일 실행 시점과 데이터 수집 상태에 따라 "
            "결과가 달라질 수 있습니다."
        )

    if metric == "top_daily_priority":
        signal = note_fields.get("signal", "확인 필요")
        score = note_fields.get("score", "확인 필요")
        priority = note.split(";")[0].strip() if note else "확인 필요"
        return (
            f"오늘 가장 먼저 확인할 종목은 {value}입니다. 우선순위는 {priority}, "
            f"신호는 {signal}, 점수는 {score}점으로 표시되었습니다. 이는 매수 추천이 "
            "아니라 리서치 우선순위입니다."
        )

    if metric == "medium_high_or_better_priority_count":
        return (
            f"Medium-High 이상으로 분류된 종목은 {value}개입니다. 오늘 집중적으로 "
            "살펴볼 후보군의 규모를 의미하며, 전부 매수 대상이라는 뜻은 아닙니다."
        )

    if metric == "high_or_very_high_risk_count":
        return (
            f"High 또는 Very High 위험도로 분류된 종목은 {value}개입니다. 우선순위가 "
            "높더라도 위험도가 높으면 진입 전 추가 확인이 필요합니다."
        )

    if metric == "type_a_priority_quality_count":
        return (
            f"우선순위와 장기 퀄리티가 함께 양호한 종목은 {value}개입니다. 다만 장기 "
            "투자 확정 신호가 아니라 추가 분석 우선순위로 해석해야 합니다."
        )

    if metric == "event_manual_review_count":
        return (
            f"최근 뉴스나 이벤트를 직접 확인해야 하는 종목은 {value}개입니다. 실적 발표, "
            "계약, 자금조달, 규제, 발사 일정 같은 이벤트 리스크를 함께 확인해야 합니다."
        )

    if metric.startswith("top_watch_"):
        ticker = metric.replace("top_watch_", "", 1) or "확인 필요"
        signal = note_fields.get("signal", "확인 필요")
        risk = note_fields.get("risk", "확인 필요")
        quality = note_fields.get("quality", "확인 필요")
        event_risk = note_fields.get("event_risk", "확인 필요")
        return (
            f"{ticker}는 상위 관심 종목 중 하나이며 현재 우선순위는 {value}입니다. "
            f"신호는 {signal}, 위험도는 {risk}, 장기 퀄리티는 {quality}, 이벤트 "
            f"리스크는 {event_risk}로 표시되었습니다. 이는 매수/매도 지시가 아니라 "
            "리서치 참고 정보입니다."
        )

    if metric == "disclaimer":
        return (
            "이 리포트는 투자 참고용 리서치 도구이며, 매수/매도 추천이나 수익률 예측이 "
            "아닙니다."
        )

    return (
        f"{metric or '확인 필요'} 항목의 값은 {value}입니다. 자세한 내용은 note를 "
        "확인해 주세요. 이 정보는 리서치 참고용이며 매수/매도 추천이 아닙니다."
    )


def add_korean_interpretations(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    interpreted_rows = []
    for row in rows:
        interpreted = dict(row)
        interpreted["korean_interpretation"] = korean_interpretation_for_row(row)
        interpreted_rows.append(interpreted)
    return interpreted_rows


def read_csv_if_exists(path: Path) -> pd.DataFrame | None:
    return pd.read_csv(path) if path.exists() else None


def run_step(step_name: str, script_path: Path, required: bool = True) -> None:
    print(f"\n=== {step_name} ===")
    try:
        result = subprocess.run(
            [sys.executable, str(script_path)],
            cwd=PROJECT_ROOT,
            text=True,
        )
        if result.returncode == 0:
            return

        message = f"{step_name} failed with exit code {result.returncode}."
        if required:
            raise RuntimeError(message)
        print(f"Warning: {message} Continuing pipeline.")
    except Exception as exc:
        if required:
            raise
        print(f"Warning: {step_name} failed: {exc}. Continuing pipeline.")


def key_metric_rows(
    priority_df: pd.DataFrame | None,
    risk_df: pd.DataFrame | None,
    quality_df: pd.DataFrame | None,
    event_df: pd.DataFrame | None,
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = [
        {
            "metric": "report_generated_at",
            "value": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "note": "Local run timestamp.",
        }
    ]

    if priority_df is not None and not priority_df.empty:
        top = priority_df.iloc[0]
        rows.append(
            {
                "metric": "top_daily_priority",
                "value": text_value(top.get("ticker")),
                "note": (
                    f"{text_value(top.get('final_priority'))}; "
                    f"signal={text_value(top.get('signal'))}; "
                    f"score={text_value(top.get('priority_score'))}"
                ),
            }
        )
        high_count = priority_df[
            priority_df["final_priority"].isin(["Very High", "High", "Medium-High"])
        ].shape[0]
        rows.append(
            {
                "metric": "medium_high_or_better_priority_count",
                "value": high_count,
                "note": "Research priority count, not a buy list.",
            }
        )

    if risk_df is not None and not risk_df.empty:
        high_risk_count = risk_df[
            risk_df["risk_label"].isin(["High", "Very High"])
        ].shape[0]
        rows.append(
            {
                "metric": "high_or_very_high_risk_count",
                "value": high_risk_count,
                "note": "Use this to separate research priority from caution level.",
            }
        )

    if quality_df is not None and not quality_df.empty:
        type_a_count = quality_df[
            quality_df["candidate_type"].astype(str).str.contains("Type A", na=False)
        ].shape[0]
        rows.append(
            {
                "metric": "type_a_priority_quality_count",
                "value": type_a_count,
                "note": "Tickers with stronger mix of priority, risk, and quality.",
            }
        )

    if event_df is not None and not event_df.empty:
        event_review_count = event_df[event_df["needs_manual_review"] == True].shape[0]
        rows.append(
            {
                "metric": "event_manual_review_count",
                "value": event_review_count,
                "note": "Tickers that need recent news/event context checked manually.",
            }
        )

    return rows


def build_executive_summary(
    priority_df: pd.DataFrame | None,
    risk_df: pd.DataFrame | None,
    quality_df: pd.DataFrame | None,
    event_df: pd.DataFrame | None,
) -> pd.DataFrame:
    rows = key_metric_rows(priority_df, risk_df, quality_df, event_df)

    if priority_df is not None and not priority_df.empty:
        risk_lookup = (
            risk_df.set_index("ticker") if risk_df is not None and "ticker" in risk_df else None
        )
        quality_lookup = (
            quality_df.set_index("ticker")
            if quality_df is not None and "ticker" in quality_df
            else None
        )
        event_lookup = (
            event_df.set_index("ticker") if event_df is not None and "ticker" in event_df else None
        )

        for _, row in priority_df.head(5).iterrows():
            ticker = text_value(row.get("ticker"))
            risk_label = ""
            quality_label = ""
            event_risk = ""
            if risk_lookup is not None and ticker in risk_lookup.index:
                risk_label = text_value(risk_lookup.loc[ticker].get("risk_label"))
            if quality_lookup is not None and ticker in quality_lookup.index:
                quality_label = text_value(
                    quality_lookup.loc[ticker].get("long_term_quality_label")
                )
            if event_lookup is not None and ticker in event_lookup.index:
                event_risk = text_value(event_lookup.loc[ticker].get("event_risk_level"))

            rows.append(
                {
                    "metric": f"top_watch_{ticker}",
                    "value": text_value(row.get("final_priority")),
                    "note": (
                        f"signal={text_value(row.get('signal'))}; "
                        f"risk={risk_label or 'n/a'}; "
                        f"quality={quality_label or 'n/a'}; "
                        f"event_risk={event_risk or 'n/a'}"
                    ),
                }
            )

    rows.append(
        {
            "metric": "disclaimer",
            "value": "research_only",
            "note": (
                "This report is not investment advice, not a price prediction, "
                "and not a trading or brokerage system."
            ),
        }
    )
    rows = add_korean_interpretations(rows)
    return pd.DataFrame(
        rows,
        columns=["metric", "value", "note", "korean_interpretation"],
    )


FINAL_SHEET_ORDER = [
    "Executive_Summary",
    "Position_Alerts",
    "Final_Daily_Priority",
    "Risk_Score",
    "Long_Term_Quality",
    "Event_Risk_Check",
    "Mega_Cap_Valuation",
    "Enhanced_Backtest_By_Signal",
    "First_Hit_By_Signal",
    "First_Hit_By_Ticker",
    "Signal_Interpretation",
]


def format_final_report() -> None:
    if not FINAL_REPORT_XLSX.exists():
        return

    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = load_workbook(FINAL_REPORT_XLSX)

    ordered_sheets = [
        workbook[sheet_name]
        for sheet_name in FINAL_SHEET_ORDER
        if sheet_name in workbook.sheetnames
    ]
    remaining_sheets = [
        workbook[sheet_name]
        for sheet_name in workbook.sheetnames
        if sheet_name not in FINAL_SHEET_ORDER
    ]
    workbook._sheets = ordered_sheets + remaining_sheets

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        headers = {
            cell.column: str(cell.value or "").strip()
            for cell in worksheet[1]
        }

        for column_index, header in headers.items():
            letter = get_column_letter(column_index)
            max_length = len(header)

            for cell in worksheet[letter]:
                if cell.row == 1:
                    continue

                value = cell.value
                if value is None:
                    continue

                text = str(value)
                max_length = max(max_length, min(len(text), 80))

                lower_header = header.lower()
                if any(
                    key in lower_header
                    for key in [
                        "return",
                        "rate",
                        "ratio",
                        "drawdown",
                        "gain",
                        "loss",
                        "margin",
                        "growth",
                        "threshold",
                        "excess",
                    ]
                ) and "pct" not in lower_header:
                    cell.number_format = "0.00%"
                elif "pct" in lower_header:
                    cell.number_format = "0.00"
                elif any(key in lower_header for key in ["price", "cash", "debt"]):
                    cell.number_format = "$#,##0.00"
                elif any(key in lower_header for key in ["date", "checked_at", "generated_at"]):
                    cell.number_format = "yyyy-mm-dd"

                if lower_header in {"note", "korean_interpretation", "position_note", "fetch_error"}:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")

            if header == "korean_interpretation":
                width = 70
            elif header in {"note", "position_note", "fetch_error"}:
                width = 50
            elif max_length > 45:
                width = 45
            else:
                width = max(12, max_length + 2)
            worksheet.column_dimensions[letter].width = width

    workbook.save(FINAL_REPORT_XLSX)


def write_final_report() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    priority_df = read_csv_if_exists(PRIORITY_CSV)
    risk_df = read_csv_if_exists(RISK_CSV)
    quality_df = read_csv_if_exists(QUALITY_CSV)
    event_df = read_csv_if_exists(EVENT_CSV)
    position_alerts_df = read_csv_if_exists(POSITION_ALERTS_CSV)
    mega_cap_valuation_df = read_csv_if_exists(MEGA_CAP_VALUATION_CSV)
    enhanced_df = read_csv_if_exists(ENHANCED_SUMMARY_CSV)
    first_hit_summary_df = read_csv_if_exists(FIRST_HIT_SUMMARY_CSV)
    first_hit_by_ticker_df = read_csv_if_exists(FIRST_HIT_BY_TICKER_CSV)
    interpretation_df = read_csv_if_exists(INTERPRETATION_CSV)
    summary_df = build_executive_summary(priority_df, risk_df, quality_df, event_df)

    with pd.ExcelWriter(FINAL_REPORT_XLSX, engine="openpyxl") as writer:
        summary_df.to_excel(writer, sheet_name="Executive_Summary", index=False)
        if priority_df is not None:
            priority_df.to_excel(writer, sheet_name="Final_Daily_Priority", index=False)
        if risk_df is not None:
            risk_df.to_excel(writer, sheet_name="Risk_Score", index=False)
        if quality_df is not None:
            quality_df.to_excel(writer, sheet_name="Long_Term_Quality", index=False)
        if event_df is not None:
            event_df.to_excel(writer, sheet_name="Event_Risk_Check", index=False)
        if position_alerts_df is not None:
            position_alerts_df.to_excel(
                writer, sheet_name="Position_Alerts", index=False
            )
        if mega_cap_valuation_df is not None:
            mega_cap_valuation_df.to_excel(
                writer, sheet_name="Mega_Cap_Valuation", index=False
            )
        if enhanced_df is not None:
            enhanced_df.to_excel(
                writer, sheet_name="Enhanced_Backtest_By_Signal", index=False
            )
        if first_hit_summary_df is not None:
            first_hit_summary_df.to_excel(
                writer, sheet_name="First_Hit_By_Signal", index=False
            )
        if first_hit_by_ticker_df is not None:
            first_hit_by_ticker_df.to_excel(
                writer, sheet_name="First_Hit_By_Ticker", index=False
            )
        if interpretation_df is not None:
            interpretation_df.to_excel(
                writer, sheet_name="Signal_Interpretation", index=False
            )

    print(f"\nWrote final report with Executive_Summary: {FINAL_REPORT_XLSX}")
    format_final_report()


def main() -> None:
    print("Starting quant_step1 daily report pipeline.")
    print("This is a research workflow only. It does not place trades.")

    for step_name, script_path, required in PIPELINE_STEPS:
        run_step(step_name, script_path, required)

    write_final_report()
    print("Done. Open output/final_daily_research_report.xlsx first.")


if __name__ == "__main__":
    main()
